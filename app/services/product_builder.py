from __future__ import annotations

import csv
import ast
import hashlib
import json
import re
import zipfile
import uuid
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from ..config import settings
from ..db import conn, dumps, loads, now_iso, row_to_dict, stable_hash
from .asset_registry import create_source_asset, link_product_asset
from .llm import chat_json
from .opportunities import analysis_is_current, get_evidence_for_opportunity, get_opportunity
from .product_lifecycle import get_product
from .validation import require_current_validation
from .paid_workflows import run_paid_once


ARTIFACT_SYSTEM = """你是数字商品交付工程师。根据已经通过真实需求分析的商品草稿，生成真正可交付的原创数字商品内容。
只允许基于输入中的真实需求证据和商品定义设计内容；不得复制第三方受版权保护内容，不得声称不存在的销量、认证、效果或用户数。
输出 JSON，必须包含 files 数组；files 中每项包含 path 和 content。至少包含 README.md 和一个真正有用的主交付文件。
如果商品明显是表格/Excel/工作簿，再输出 workbooks 数组，每个工作簿包含 path 和 sheets；sheet 包含 name、columns、rows。
文本文件只允许 md/txt/csv/json/py/js/html/css/yaml/yml。内容必须能直接被买家使用，而不是只给一个目录或空壳说明。"""


_ALLOWED_TEXT_EXTENSIONS = {".md", ".txt", ".csv", ".json", ".py", ".js", ".html", ".css", ".yaml", ".yml"}
_MAX_FILES = 12
_MAX_FILE_CHARS = 120_000
_MAX_ROWS_PER_SHEET = 3000
_MAX_COLS_PER_SHEET = 80

_FORMAT_EXTENSIONS = {
    "xlsx": {".xlsx"},
    "csv": {".csv"},
    "markdown": {".md"},
    "txt": {".txt"},
    "json": {".json"},
    "python": {".py"},
    "javascript": {".js"},
    "html": {".html"},
    "css": {".css"},
    "yaml": {".yaml", ".yml"},
}


def _safe_relative_path(value: Any) -> Path:
    raw = str(value or "").replace("\\", "/")
    # Enforce Windows-safe paths on every host, including Linux CI.
    parts = raw.split("/")
    reserved = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
    if (not raw or raw.startswith((".", "/")) or any(
        not part or part in {".", ".."} or part.endswith((".", " "))
        or any(char in part for char in ':<>"|?*')
        or any(ord(char) < 32 for char in part)
        or part.split(".")[0].upper() in reserved for part in parts
    )):
        raise ValueError(f"unsafe artifact path: {value!r}")
    path = Path(raw)
    if len(path.parts) > 5:
        raise ValueError(f"artifact path too deep: {raw}")
    return path


def _write_text_file(root: Path, spec: dict[str, Any]) -> str:
    rel = _safe_relative_path(spec.get("path"))
    if rel.suffix.lower() not in _ALLOWED_TEXT_EXTENSIONS:
        raise ValueError(f"unsupported artifact extension: {rel.suffix}")
    content = spec.get("content")
    if not isinstance(content, str) or not content.strip():
        raise ValueError(f"artifact file is empty: {rel}")
    if len(content) > _MAX_FILE_CHARS:
        raise ValueError(f"artifact file too large: {rel}")
    if rel.suffix.lower() == ".json":
        json.loads(content)
    elif rel.suffix.lower() == ".py":
        ast.parse(content, filename=str(rel))
    elif rel.suffix.lower() == ".csv":
        rows = list(csv.reader(content.splitlines(), strict=True))
        if len(rows) < 2 or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
            raise ValueError(f"CSV must contain a header and consistent data rows: {rel}")
    target = root / rel
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content, encoding="utf-8", newline="\n")
    return rel.as_posix()


def _write_workbook(root: Path, spec: dict[str, Any]) -> str:
    rel = _safe_relative_path(spec.get("path"))
    if rel.suffix.lower() != ".xlsx":
        raise ValueError("workbook path must end with .xlsx")
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError(f"workbook has no sheets: {rel}")

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    if len(sheets) > 12:
        raise ValueError("workbook has too many sheets")
    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            raise ValueError("workbook sheet must be an object")
        name = str(sheet.get("name") or f"Sheet{index + 1}")[:31]
        ws = wb.create_sheet(title=name)
        columns = sheet.get("columns") or []
        rows = sheet.get("rows") or []
        if not isinstance(columns, list) or not columns:
            raise ValueError(f"workbook sheet missing columns: {name}")
        if not isinstance(rows, list):
            raise ValueError(f"workbook rows must be an array: {name}")
        if not rows or len(rows) > _MAX_ROWS_PER_SHEET or len(columns) > _MAX_COLS_PER_SHEET:
            raise ValueError(f"workbook needs substantive rows within size limits: {name}")
        columns = [str(v) for v in columns]
        ws.append(columns)
        for row in rows[:_MAX_ROWS_PER_SHEET]:
            if isinstance(row, dict):
                values = [row.get(column) for column in columns]
            elif isinstance(row, list):
                if len(row) != len(columns):
                    raise ValueError(f"workbook row width does not match columns: {name}")
                values = list(row)
            else:
                raise ValueError(f"unsupported workbook row in {name}")
            ws.append(values)
        if not any(cell.value is not None and str(cell.value).strip() for row in ws.iter_rows(min_row=2) for cell in row):
            raise ValueError(f"workbook has no substantive cells: {name}")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    target = root / rel
    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return rel.as_posix()


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _validate_manifest(value: dict[str, Any], expected_format: str) -> dict[str, Any]:
    files = value.get("files")
    if not isinstance(files, list) or not files:
        raise ValueError("artifact generator must return files including README.md")
    if len(files) > _MAX_FILES:
        raise ValueError("artifact generator returned too many files")
    names = [str(item.get("path") or "") for item in files if isinstance(item, dict)]
    if "readme.md" not in {name.lower() for name in names}:
        raise ValueError("artifact generator must include README.md")
    workbooks = value.get("workbooks") or []
    if not isinstance(workbooks, list):
        raise ValueError("workbooks must be an array")
    specs = files + workbooks
    if len(specs) > _MAX_FILES or any(not isinstance(item, dict) for item in specs):
        raise ValueError("artifact specs must be objects within file limits")
    paths = [_safe_relative_path(item.get("path")).as_posix().casefold() for item in specs]
    if len(set(paths)) != len(paths) or "manifest.json" in paths:
        raise ValueError("duplicate or reserved artifact path")
    if len(files) + len(workbooks) < 2:
        raise ValueError("artifact generator must return README plus at least one primary deliverable")

    allowed = _FORMAT_EXTENSIONS.get(expected_format)
    if not allowed:
        raise ValueError(f"unsupported expected artifact format: {expected_format}")
    candidate_paths = names + [str(item.get("path") or "") for item in workbooks if isinstance(item, dict)]
    if not any(Path(path).suffix.lower() in allowed and Path(path).name.lower() != "readme.md" for path in candidate_paths):
        raise ValueError(f"artifact generator did not produce required {expected_format} deliverable")
    return value


def _build_product_artifact(product_id: int) -> dict[str, Any]:
    product = get_product(product_id)
    if not product:
        raise KeyError("product not found")
    opportunity = get_opportunity(int(product["opportunity_id"]))
    if not opportunity:
        raise KeyError("opportunity not found")
    validation = require_current_validation(int(product["opportunity_id"]))
    if not analysis_is_current(opportunity):
        raise ValueError("opportunity analysis is stale; rebuild before producing deliverables")

    artifact_format = str(product.get("artifact_format") or "").strip().lower()
    if artifact_format not in _FORMAT_EXTENSIONS:
        raise ValueError("product has no supported artifact_format; rerun product generation from a current BUILD card")
    evidence = get_evidence_for_opportunity(int(product["opportunity_id"]), 30)
    analysis = loads(opportunity.get("analysis_json"), {}) or {}
    if str(analysis.get("decision") or "").upper() != "BUILD":
        raise ValueError("current opportunity decision is not BUILD")
    payload = {
        "product": {
            "id": product_id,
            "sku": product.get("sku"),
            "name": product.get("name"),
            "category": product.get("category"),
            "audience": product.get("audience"),
            "title": product.get("title"),
            "description": product.get("description"),
            "deliverables": product.get("deliverables") or [],
            "faq": product.get("faq") or [],
            "compliance_notes": product.get("compliance_notes"),
            "artifact_format": artifact_format,
        },
        "opportunity": {
            "id": opportunity.get("id"),
            "title": opportunity.get("title"),
            "normalized_need": opportunity.get("normalized_need"),
            "suggested_product": opportunity.get("suggested_product"),
        },
        "analysis": {
            "decision": analysis.get("decision"),
            "reasoning": analysis.get("reasoning"),
            "pain_points": analysis.get("pain_points") or [],
            "target_users": analysis.get("target_users") or [],
            "recommended_products": analysis.get("recommended_products") or [],
        },
        "evidence": [
            {
                "id": item.get("id"),
                "platform": item.get("platform"),
                "title": item.get("title"),
                "content": str(item.get("content") or "")[:700],
                "source_url": item.get("source_url"),
            }
            for item in evidence[:20]
        ],
        "output_schema": {
            "files": [{"path": "README.md", "content": "complete buyer-facing instructions"}],
            "workbooks": [{"path": "template.xlsx", "sheets": [{"name": "Sheet1", "columns": ["column"], "rows": [["example"]]}]}],
            "quality_notes": ["string"],
        },
    }

    generated = chat_json(
        operation="product_artifact_build",
        entity_type="product",
        entity_id=product_id,
        system=ARTIFACT_SYSTEM,
        payload=payload,
        max_new_tokens=max(2200, settings.qiniu_json_retry_max_tokens),
    )
    if require_current_validation(int(product["opportunity_id"]))["snapshot_hash"] != validation["snapshot_hash"]:
        raise ValueError("demand evidence changed during artifact generation")
    manifest = _validate_manifest(generated, artifact_format)

    base = Path(settings.artifact_dir).expanduser().resolve()
    root = base / f"product_{product_id}"
    root.mkdir(parents=True, exist_ok=True)
    build_stamp = re.sub(r"[^0-9]", "", now_iso())[:14]
    work = root / f"build_{build_stamp}_{uuid.uuid4().hex}"
    work.mkdir(parents=True, exist_ok=False)

    written: list[str] = []
    for spec in manifest["files"]:
        if not isinstance(spec, dict):
            raise ValueError("artifact file spec must be an object")
        written.append(_write_text_file(work, spec))
    for spec in manifest.get("workbooks") or []:
        if not isinstance(spec, dict):
            raise ValueError("workbook spec must be an object")
        written.append(_write_workbook(work, spec))

    if len(written) < 2:
        raise ValueError("artifact build produced fewer than two usable files")

    manifest_path = work / "manifest.json"
    manifest_payload = {
        "product_id": product_id,
        "sku": product.get("sku"),
        "artifact_format": artifact_format,
        "generated_at": now_iso(),
        "files": written,
        "quality_notes": manifest.get("quality_notes") or [],
        "evidence_ids": [item.get("id") for item in evidence],
        "policy": "original generated deliverable grounded in stored real demand evidence",
    }
    manifest_path.write_text(json.dumps(manifest_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    # Each build remains immutable; rebuilding cannot overwrite an approved package.
    zip_path = work / f"product-{product_id}.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(work.rglob("*")):
            if path.is_file():
                if path == zip_path:
                    continue
                archive.write(path, path.relative_to(work).as_posix())
    checksum = _sha256(zip_path)

    asset = create_source_asset(
        {
            "source_type": "owned",
            "title": f"{product.get('name') or product.get('sku')} deliverable package",
            "local_ref": str(zip_path),
            "checksum": checksum,
            "ownership_type": "self_owned",
            "license_name": "Original generated product",
            "commercial_use_allowed": True,
            "redistribution_allowed": True,
            "license_evidence": {
                "generator": "digital-goods-agent-engine",
                "product_id": product_id,
                "build_manifest": str(manifest_path),
                "generated_at": manifest_payload["generated_at"],
                "evidence_ids": manifest_payload["evidence_ids"],
            },
        }
    )
    link_product_asset(product_id, int(asset["id"]), "deliverable")

    created_at = now_iso()
    with conn() as c:
        cur = c.execute(
            """INSERT INTO product_builds(
                   product_id,status,bundle_path,checksum,manifest_json,source_asset_id,error,created_at,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?)""",
            (
                product_id,
                "SUCCESS",
                str(zip_path),
                checksum,
                dumps(manifest_payload),
                int(asset["id"]),
                None,
                created_at,
                created_at,
            ),
        )
        build_id = int(cur.lastrowid)
    return {
        "build_id": build_id,
        "product_id": product_id,
        "status": "SUCCESS",
        "bundle_path": str(zip_path),
        "checksum": checksum,
        "source_asset_id": int(asset["id"]),
        "files": written,
        "manifest": manifest_payload,
    }


def list_product_builds(product_id: int, limit: int = 20) -> list[dict[str, Any]]:
    with conn() as c:
        rows = c.execute(
            "SELECT * FROM product_builds WHERE product_id=? ORDER BY id DESC LIMIT ?",
            (product_id, max(1, min(limit, 100))),
        ).fetchall()
    result = []
    for row in rows:
        item = row_to_dict(row) or {}
        item["manifest"] = loads(item.pop("manifest_json", None), {}) or {}
        result.append(item)
    return result


def latest_successful_bundle(product_id: int) -> dict[str, Any]:
    with conn() as c:
        row = c.execute(
            "SELECT * FROM product_builds WHERE product_id=? AND status='SUCCESS' ORDER BY id DESC LIMIT 1",
            (product_id,),
        ).fetchone()
    if not row:
        raise ValueError("product has no successfully generated deliverable package")
    item = row_to_dict(row) or {}
    path = Path(str(item.get("bundle_path") or "")).resolve()
    if not path.is_file():
        raise ValueError("generated deliverable package is missing from disk")
    checksum = _sha256(path)
    if checksum != item.get("checksum"):
        raise ValueError("generated deliverable package checksum mismatch")
    item["manifest"] = loads(item.pop("manifest_json", None), {}) or {}
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise ValueError("generated deliverable package has corrupt members")
            files = item["manifest"].get("files") or []
            if not files or not set(files).issubset(archive.namelist()):
                raise ValueError("generated deliverable package is missing manifest files")
    except zipfile.BadZipFile as exc:
        raise ValueError("generated deliverable package is not a valid ZIP") from exc
    return item


def build_product_artifact(product_id: int, *, force: bool = False) -> dict[str, Any]:
    product = get_product(product_id)
    if not product:
        raise KeyError("product not found")
    opportunity = get_opportunity(int(product["opportunity_id"]))
    validation = require_current_validation(int(product["opportunity_id"]))
    if not opportunity or not analysis_is_current(opportunity):
        raise ValueError("opportunity analysis is stale; rebuild before producing deliverables")
    fields = {key:product.get(key) for key in ("sku","name","category","audience","title","description","deliverables","faq","compliance_notes","artifact_format")}
    snapshot = stable_hash("artifact-dga001-v1",dumps(fields),opportunity.get("analysis_json"),validation["snapshot_hash"])
    result = run_paid_once("product_artifact",product_id,snapshot,lambda:_build_product_artifact(product_id),force=force)
    if result.get("idempotent"):
        path=Path(result["bundle_path"])
        if not path.is_file() or _sha256(path) != result["checksum"]:
            raise ValueError("cached artifact is missing or changed; inspect the build and explicitly force a rebuild")
    return result
